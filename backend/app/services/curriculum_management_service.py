"""Persistent curriculum imports, graph-constrained coverage and optimization runs."""

from __future__ import annotations

import hashlib
import json
import re
import uuid
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any

from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.enums import DataFlag, GraphStatus
from app.core.errors import ConflictError, NotFoundError, ValidationError
from app.models.competency import CompetencyGraph, CompetencyNode
from app.models.curriculum import (
    CourseSkillCoverage, CurriculumAnalysis, CurriculumCourse, CurriculumDraftCourse,
    CurriculumImportBatch, CurriculumPlan, OptimizationRun, OptimizationSuggestion,
    SkillCoverage, SuggestionEvidence,
)
from app.models.job_market import JobPosting, JobPostingSkill, SkillDemandSnapshot
from app.models.ontology import Job, Skill
from app.rag.chunker import chunk_pages
from app.rag.loader import load_document

MAX_BYTES = 10 * 1024 * 1024
COURSE_FIELDS = ("objectives", "description", "teaching_content", "knowledge_points", "practical_content", "learning_outcomes")
LABELS = {
    "课程目标": "objectives", "目标": "objectives", "课程简介": "description", "简介": "description",
    "教学内容": "teaching_content", "主要内容": "teaching_content", "知识点": "knowledge_points",
    "实践内容": "practical_content", "实训内容": "practical_content", "学习成果": "learning_outcomes",
    "课程代码": "course_code", "课程编号": "course_code", "课程类别": "category", "总学时": "total_hours",
}


def _evidence(value: str, chunk_id: str, page: str | None) -> dict[str, Any]:
    return {"chunk_id": chunk_id, "page": page, "quote": value}


_COURSE_CODE_LINE = re.compile(r"^\s*(\d{8,12})\s*$")
_COURSE_HOURS_LINE = re.compile(r"^\s*(\d+(?:\.\d+)?)\s+(\d+)\b")


def _table_courses(document_text: str) -> list[tuple[dict, dict]]:
    """Extract course rows from common PDF curriculum tables.

    PDF text extraction usually flattens a table into a course-code line,
    Chinese course name, English name and a numeric credit/hours line.  This
    parser intentionally requires both a code and a numeric row so ordinary
    prose headings are not mistaken for courses.  It is a deterministic
    fallback for documents that do not use ``课程名称：`` labels.
    """
    lines = document_text.splitlines()
    results: list[tuple[dict, dict]] = []
    # Curriculum PDFs often repeat the same course in multiple semester or
    # elective tables.  The persisted course model is unique by plan/name, so
    # retain the first complete row for each normalized course name.
    seen: set[str] = set()

    def clean_name(value: str) -> str:
        value = re.sub(r"\s+", " ", value).strip()
        # The extractor commonly places the English translation on the same
        # line as a Chinese name (for example ``...基础 The``). Keep useful
        # level markers such as ``A I`` / ``II`` because they distinguish
        # otherwise identical courses.
        match = re.match(r"^(.+?[\u3400-\u9fff])(?:\s+((?:[AB]|I{1,3}|IV|V)(?:\s+(?:I{1,3}|IV|V))?))?$", value)
        if match:
            return f"{match.group(1)} {match.group(2)}".strip() if match.group(2) else match.group(1).strip()
        return value

    for index, line in enumerate(lines):
        code_match = _COURSE_CODE_LINE.fullmatch(line)
        if not code_match:
            continue
        code = code_match.group(1)
        name: str | None = None
        hours: int | None = None
        end = min(index + 8, len(lines))
        for cursor in range(index + 1, end):
            candidate = lines[cursor].strip()
            if not candidate:
                continue
            if _COURSE_CODE_LINE.fullmatch(candidate):
                break
            hours_match = _COURSE_HOURS_LINE.match(candidate)
            if hours_match:
                hours = int(hours_match.group(2))
                end = cursor + 1
                break
            if name is None and any("\u3400" <= char <= "\u9fff" for char in candidate):
                name = clean_name(candidate)
        if not name or hours is None or len(name) < 2:
            continue
        key = re.sub(r"\s+", "", name).casefold()
        if key in seen:
            continue
        seen.add(key)
        quote = "\n".join(lines[index:end]).strip()
        evidence = {
            "name": _evidence(name, f"table:{len(results) + 1}", None),
            "course_code": _evidence(code, f"table:{len(results) + 1}", None),
            "total_hours": _evidence(str(hours), f"table:{len(results) + 1}", None),
        }
        results.append((
            {"name": name, "course_code": code, "total_hours": hours, "source_quote": quote},
            evidence,
        ))
    return results


def _structured_courses(filename: str, content: bytes, stored_path: Path) -> tuple[str, list[tuple[dict, dict]]]:
    suffix = stored_path.suffix.lower()
    if suffix == ".json":
        try:
            payload = json.loads(content.decode("utf-8-sig"))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise ValidationError("课程 JSON 格式或编码不正确。") from exc
        items = payload.get("courses") if isinstance(payload, dict) else payload
        if not isinstance(items, list) or not all(isinstance(item, dict) for item in items):
            raise ValidationError("课程 JSON 必须是课程数组或包含 courses 数组。")
        document_text = json.dumps(payload, ensure_ascii=False)
        result: list[tuple[dict, dict]] = []
        for index, item in enumerate(items, 1):
            name = str(item.get("name") or item.get("course_name") or "").strip()
            if not name:
                result.append(({"name": ""}, {}))
                continue
            data = {key: item.get(key) for key in ("id", "course_code", "name", "category", "total_hours", *COURSE_FIELDS)}
            data["name"] = name
            if item.get("source_quote"):
                data["source_quote"] = item["source_quote"]
            evidence = {key: _evidence(str(value), f"json:{index}", None) for key, value in data.items() if value not in (None, "")}
            result.append((data, evidence))
        return document_text, result

    pages = load_document(stored_path)
    chunks = chunk_pages(pages)
    document_text = "\n".join(page.text for page in pages)
    if not document_text.strip():
        if suffix == ".pdf":
            raise ValidationError("PDF 未提取到文字，可能是扫描件；请上传文字版或 OCR 版。")
        raise ValidationError("课程文件未提取到文字。")
    # Course blocks start at an explicit course-name label. This is deliberately
    # conservative: an unlabelled line is not silently invented as a course.
    pattern = re.compile(r"(?m)^\s*(?:课程名称|课程名)\s*[:：\t]\s*(.+?)\s*$")
    matches = list(pattern.finditer(document_text))
    results: list[tuple[dict, dict]] = []
    for pos, match in enumerate(matches):
        end = matches[pos + 1].start() if pos + 1 < len(matches) else len(document_text)
        block = document_text[match.start():end].strip()
        name = match.group(1).strip()
        data: dict[str, Any] = {"name": name}
        field_ev = {"name": _evidence(match.group(0).strip(), f"text:{pos + 1}", None)}
        for line in block.splitlines()[1:]:
            field_match = re.match(r"\s*([^:：\t]{1,10})\s*[:：\t]\s*(.+)\s*$", line)
            if not field_match:
                continue
            field = LABELS.get(field_match.group(1).strip())
            if not field:
                continue
            value = field_match.group(2).strip()
            data[field] = int(value) if field == "total_hours" and value.isdigit() else value
            field_ev[field] = _evidence(line.strip(), f"text:{pos + 1}", None)
        data["source_quote"] = block
        results.append((data, field_ev))
    if not results:
        results = _table_courses(document_text)
    if not results:
        # Retain a reviewable draft rather than pretending extraction succeeded.
        excerpt = document_text.strip()[:2000]
        results.append(({"name": Path(filename).stem, "description": excerpt, "source_quote": excerpt, "_needs_ai_structure": True}, {"description": _evidence(excerpt, "text:1", None)}))
    return document_text, results


class CurriculumManagementService:
    def __init__(self, db: Session) -> None:
        self.db = db

    def create_import(self, *, filename: str, content: bytes, source_name: str, source_url: str | None, license_note: str | None) -> CurriculumImportBatch:
        if not content:
            raise ValidationError("上传文件为空。")
        if len(content) > MAX_BYTES:
            raise ValidationError("课程文件超过 10MB 限制。")
        suffix = Path(filename).suffix.lower()
        if suffix not in {".pdf", ".docx", ".txt", ".json"}:
            raise ValidationError("课程导入仅支持 PDF、DOCX、TXT 和结构化 JSON。")
        file_hash = hashlib.sha256(content).hexdigest()
        storage = settings.data_dir / "curriculum_uploads"
        storage.mkdir(parents=True, exist_ok=True)
        batch_id = f"cib_{uuid.uuid4().hex}"
        path = storage / f"{batch_id}{suffix}"
        path.write_bytes(content)
        batch = CurriculumImportBatch(
            id=batch_id, filename=filename, file_hash=file_hash, storage_path=str(path),
            source_name=source_name, source_url=source_url, license_note=license_note,
        )
        self.db.add(batch)
        self.db.flush()
        try:
            self.parse_import(batch)
        except Exception:
            # A failed parse remains addressable and can be retried after replacing
            # the parser/dependency; the uploaded source and hash are retained.
            pass
        return batch

    def get_import(self, batch_id: str) -> CurriculumImportBatch:
        batch = self.db.get(CurriculumImportBatch, batch_id)
        if not batch:
            raise NotFoundError(f"未找到课程导入批次：{batch_id}")
        return batch

    def drafts(self, batch_id: str) -> list[CurriculumDraftCourse]:
        return list(self.db.execute(select(CurriculumDraftCourse).where(CurriculumDraftCourse.batch_id == batch_id).order_by(CurriculumDraftCourse.row_number)).scalars())

    def parse_import(self, batch: CurriculumImportBatch) -> CurriculumImportBatch:
        if batch.status == "confirmed":
            raise ConflictError("已确认的课程批次不能重新解析。")
        self.db.execute(delete(CurriculumDraftCourse).where(CurriculumDraftCourse.batch_id == batch.id))
        try:
            content = Path(batch.storage_path).read_bytes()
            text, courses = _structured_courses(batch.filename, content, Path(batch.storage_path))
            batch.document_text, batch.error_message, batch.status = text, None, "parsed"
            needs_ai = len(courses) == 1 and bool(courses[0][0].get("_needs_ai_structure"))
            if needs_ai and settings.llm_api_key:
                try:
                    from app.agents.curriculum_structure import CurriculumStructureAgent, CurriculumStructureInput
                    envelope = CurriculumStructureAgent().run(self.db, CurriculumStructureInput(filename=batch.filename, document_text=text))
                    extracted: list[tuple[dict, dict]] = []
                    for index, item in enumerate(envelope.result.courses, 1):
                        data = item.model_dump(exclude={"evidence_quotes"})
                        nonempty = [key for key, value in data.items() if value not in (None, "")]
                        for key in nonempty:
                            quote = item.evidence_quotes.get(key, "").strip()
                            if not quote or quote not in text:
                                raise ValidationError(f"课程结构 Agent 的 {item.name}/{key} 缺少可验证连续引文。")
                        evidence = {key: _evidence(item.evidence_quotes[key], f"ai:{index}", None) for key in nonempty}
                        evidence["_generation"] = {"llm_run_id": envelope.llm_run_id, "confidence": envelope.confidence, "reason": envelope.reasoning_summary}
                        data["source_quote"] = max(item.evidence_quotes.values(), key=len)
                        extracted.append((data, evidence))
                    if not extracted:
                        raise ValidationError("课程结构 Agent 未识别到可核验课程。")
                    courses = extracted
                except Exception:
                    batch.error_message = "AI 课程结构化失败，已保留可人工复核的原文草稿；可配置模型后重试或确认后手工编辑。"
            elif needs_ai:
                batch.error_message = "未配置真实 .env 中的 LLM_API_KEY，已保留可人工复核的原文草稿；明确标签的课程仍会由规则正常解析。"
            for index, (data, evidence) in enumerate(courses, 1):
                data.pop("_needs_ai_structure", None)
                errors = [] if str(data.get("name") or "").strip() else ["课程名称为必填字段"]
                self.db.add(CurriculumDraftCourse(batch_id=batch.id, row_number=index, structured_data=data, field_evidence=evidence, status="valid" if not errors else "invalid", errors=errors))
        except Exception as exc:
            batch.status, batch.error_message = "failed", str(exc)
            self.db.flush()
            raise
        self.db.flush()
        return batch

    def confirm_import(self, batch_id: str, *, plan_id: str, name: str, profession: str | None, version: str | None, is_partial: bool) -> CurriculumPlan:
        batch = self.get_import(batch_id)
        if batch.status == "confirmed":
            plan = self.db.get(CurriculumPlan, batch.confirmed_plan_id)
            if plan:
                return plan
        if batch.status != "parsed":
            raise ConflictError("课程批次尚未成功解析。")
        if self.db.get(CurriculumPlan, plan_id):
            raise ConflictError(f"培养方案 ID 已存在：{plan_id}")
        drafts = self.drafts(batch_id)
        if not drafts:
            raise ValidationError("课程草稿为空，不能确认。")
        # Older batches may have been staged before the PDF table parser was
        # added and can contain repeated names from several tables.  The
        # persisted course model intentionally keeps one row per plan/name;
        # retain the first valid draft so confirmation remains recoverable.
        unique_drafts: list[CurriculumDraftCourse] = []
        seen_names: set[str] = set()
        for draft in drafts:
            value = str(draft.structured_data.get("name") or "").strip()
            key = re.sub(r"\s+", "", value).casefold()
            if key and key in seen_names:
                continue
            if key:
                seen_names.add(key)
            unique_drafts.append(draft)
        if any(item.status != "valid" for item in unique_drafts):
            raise ValidationError("课程草稿存在校验错误，不能确认。")
        plan = CurriculumPlan(id=plan_id, name=name, profession=profession, version=version, source_doc_id=batch.id, source_name=batch.source_name, source_url=batch.source_url, license_note=batch.license_note, is_partial=is_partial)
        self.db.add(plan)
        self.db.flush()
        course_ids: dict[str, str] = {}
        for draft in unique_drafts:
            data = draft.structured_data
            quote = str(data.get("source_quote") or data.get("description") or data.get("name"))
            if quote not in batch.document_text:
                # JSON values appear in the serialized original and are still mechanically verifiable.
                candidates = [str(v) for v in data.values() if v not in (None, "") and str(v) in batch.document_text]
                quote = max(candidates, key=len) if candidates else str(data["name"])
            course_id = str(data.get("id") or f"{plan_id}.course.{draft.row_number}")
            self.db.add(CurriculumCourse(
                id=course_id, plan_id=plan_id, course_code=data.get("course_code"), name=data["name"],
                category=data.get("category"), total_hours=data.get("total_hours"),
                source_chunk_id=f"{batch.id}:{draft.row_number}", source_page=None, source_quote=quote,
                field_evidence=draft.field_evidence, **{field: data.get(field) for field in COURSE_FIELDS},
            ))
            course_ids[str(data.get("id") or "")] = course_id
        # CourseSkillCoverage has no ORM relationship to communicate insert
        # ordering, so make the referenced imported courses visible first.
        self.db.flush()
        # Structured JSON may carry reviewed coverage rows alongside courses.
        # Preserve them as teacher/rule mappings instead of silently dropping
        # the evidence during confirmation.
        try:
            payload = json.loads(Path(batch.storage_path).read_text(encoding="utf-8-sig")) if Path(batch.storage_path).suffix.lower() == ".json" else {}
            for item in (payload.get("coverage", []) if isinstance(payload, dict) else []):
                course_id = course_ids.get(str(item.get("course_id")))
                skill_code = str(item.get("skill_code") or "")
                if not course_id:
                    raise ValidationError(f"课程技能映射引用了不存在的课程：{item.get('course_id')}")
                if not self.db.get(Skill, skill_code):
                    raise ValidationError(f"课程技能映射必须使用规范技能代码：{skill_code or '(空)'}")
                status = str(item.get("coverage_status") or "").strip().lower()
                if status not in {"covered", "partial", "uncovered"}:
                    strength_hint = int(item.get("coverage_strength") or 0)
                    status = "covered" if strength_hint >= 2 else "partial" if strength_hint == 1 else "uncovered"
                strength = {"uncovered": 0, "partial": 1, "covered": 2}[status]
                course = self.db.get(CurriculumCourse, course_id)
                quote = str(item.get("evidence_quote") or "").strip()
                if status != "uncovered" and (not quote or quote not in self.course_corpus(course)):
                    raise ValidationError(f"{course.name} / {skill_code} 的证据不是课程原文中的连续子串。")
                self.db.add(CourseSkillCoverage(course_id=course_id, skill_code=skill_code, coverage_strength=strength, coverage_status=status, origin="teacher" if item.get("teacher_confirmed") else "rule", confidence=item.get("confidence"), reason=item.get("reason"), teacher_confirmed=bool(item.get("teacher_confirmed")), edited_by=item.get("edited_by"), source_chunk_id=str(item.get("source_chunk_id") or course.source_chunk_id), source_page=item.get("source_page") or course.source_page, evidence_quote=quote))
        except (OSError, json.JSONDecodeError):
            pass
        batch.status, batch.confirmed_plan_id, batch.confirmed_at = "confirmed", plan_id, datetime.now(timezone.utc)
        self.db.flush()
        return plan

    def add_course(self, plan_id: str, data: dict[str, Any]) -> CurriculumCourse:
        if not self.db.get(CurriculumPlan, plan_id):
            raise NotFoundError(f"未找到培养方案：{plan_id}")
        course_id = data.pop("id", None) or f"course_{uuid.uuid4().hex}"
        course = CurriculumCourse(id=course_id, plan_id=plan_id, field_evidence={}, **data)
        self.db.add(course); self.db.flush()
        return course

    @staticmethod
    def course_corpus(course: CurriculumCourse) -> str:
        # Prefer semantically rich structured fields; source_quote is retained
        # as a last-resort trace to the original course block.
        return "\n".join(str(value) for value in (*(getattr(course, field) for field in COURSE_FIELDS), course.source_quote) if value)

    def upsert_mapping(self, course_id: str, data: dict[str, Any], mapping_id: int | None = None) -> CourseSkillCoverage:
        course = self.db.get(CurriculumCourse, course_id)
        if not course:
            raise NotFoundError(f"未找到课程：{course_id}")
        if not self.db.get(Skill, data["skill_code"]):
            raise ValidationError("技能必须来自规范 Skill.skill_code。")
        evidence_quote = str(data.get("evidence_quote") or "").strip()
        if data["coverage_status"] != "uncovered":
            if not evidence_quote:
                raise ValidationError("已覆盖或部分覆盖必须提供课程原文证据。")
            if evidence_quote not in self.course_corpus(course):
                raise ValidationError("证据引文不是课程原文中的连续子串。")
        strength = {"uncovered": 0, "partial": 1, "covered": 2}[data["coverage_status"]]
        item = self.db.get(CourseSkillCoverage, mapping_id) if mapping_id else self.db.execute(select(CourseSkillCoverage).where(CourseSkillCoverage.course_id == course_id, CourseSkillCoverage.skill_code == data["skill_code"])).scalar_one_or_none()
        if item is None:
            item = CourseSkillCoverage(course_id=course_id, skill_code=data["skill_code"], coverage_strength=strength, origin="teacher" if data.get("edited_by") else "rule", **{k: v for k, v in data.items() if k != "coverage_status"})
            self.db.add(item)
        else:
            for key, value in data.items(): setattr(item, key, value)
            item.coverage_strength, item.origin = strength, "teacher"
        item.coverage_status = data["coverage_status"]
        self.db.flush()
        self.refresh_plan_analyses(course.plan_id)
        return item

    def refresh_plan_analyses(self, plan_id: str) -> None:
        analyses = list(self.db.execute(select(CurriculumAnalysis).where(CurriculumAnalysis.plan_id == plan_id)).scalars())
        for analysis in analyses:
            self.build_analysis(analysis.job_id, analysis.plan_id, analysis.graph_id, existing=analysis)
        if analyses:
            ids = [a.id for a in analyses]
            for run in self.db.execute(select(OptimizationRun).where(OptimizationRun.analysis_id.in_(ids))).scalars():
                run.is_stale = True

    def build_analysis(self, job_id: str, plan_id: str, graph_id: str, *, existing: CurriculumAnalysis | None = None) -> CurriculumAnalysis:
        if not self.db.get(Job, job_id): raise NotFoundError(f"未找到岗位：{job_id}")
        if not self.db.get(CurriculumPlan, plan_id): raise NotFoundError(f"未找到培养方案：{plan_id}")
        graph = self.db.get(CompetencyGraph, graph_id)
        if not graph or graph.job_id != job_id or graph.status is not GraphStatus.APPROVED:
            raise ValidationError("正式对标必须绑定该岗位已审核的能力图谱。")
        snapshot_time = self.db.execute(select(func.max(SkillDemandSnapshot.computed_at)).where(SkillDemandSnapshot.job_id == job_id)).scalar_one()
        if snapshot_time is None:
            raise ValidationError("该岗位没有需求快照，请先执行岗位分析。")
        version = snapshot_time.isoformat()
        analysis = existing or CurriculumAnalysis(id=f"ca_{uuid.uuid4().hex}", job_id=job_id, plan_id=plan_id, graph_id=graph_id, snapshot_version=version)
        if not existing: self.db.add(analysis); self.db.flush()
        analysis.snapshot_version = version
        self.db.execute(delete(SkillCoverage).where(SkillCoverage.analysis_id == analysis.id))
        nodes = list(self.db.execute(select(CompetencyNode).where(CompetencyNode.graph_id == graph_id, CompetencyNode.skill_code.is_not(None))).scalars())
        mastery: dict[str, int] = {}
        for node in nodes: mastery[node.skill_code] = max(mastery.get(node.skill_code, 1), node.mastery_level or 1)
        courses = list(self.db.execute(select(CurriculumCourse).where(CurriculumCourse.plan_id == plan_id)).scalars())
        skill_codes = set(mastery)
        skills = {s.skill_code: s for s in self.db.execute(select(Skill).where(Skill.skill_code.in_(skill_codes))).scalars()} if skill_codes else {}
        # Deterministic rule mapping constrained to graph skill codes.
        for course in courses:
            corpus = self.course_corpus(course)
            for code, skill in skills.items():
                current = self.db.execute(select(CourseSkillCoverage).where(CourseSkillCoverage.course_id == course.id, CourseSkillCoverage.skill_code == code)).scalar_one_or_none()
                if current and (current.teacher_confirmed or current.origin == "teacher"):
                    continue
                terms = sorted({t for t in (skill.name_zh, skill.name_en, skill.skill_code, *skill.aliases) if t and len(t) >= 2}, key=len, reverse=True)
                match = next((re.search(re.escape(term), corpus, re.I) for term in terms if re.search(re.escape(term), corpus, re.I)), None)
                if match:
                    left = max(corpus.rfind(mark, 0, match.start()) for mark in "。！？；\n") + 1
                    rights = [corpus.find(mark, match.end()) for mark in "。！？；\n"]
                    right = min((v for v in rights if v >= 0), default=len(corpus)) + 1
                    quote = corpus[left:right].strip()
                    if current:
                        current.coverage_strength, current.coverage_status = 1, "partial"
                        current.confidence, current.reason = 1.0, "规范技能名或别名在课程原文中精确命中"
                        current.source_chunk_id, current.source_page, current.evidence_quote = course.source_chunk_id, course.source_page, quote
                    else:
                        self.db.add(CourseSkillCoverage(course_id=course.id, skill_code=code, coverage_strength=1, coverage_status="partial", origin="rule", confidence=1.0, reason="规范技能名或别名在课程原文中精确命中", source_chunk_id=course.source_chunk_id, source_page=course.source_page, evidence_quote=quote))
                elif current:
                    self.db.delete(current)
        self.db.flush()
        mappings = list(self.db.execute(select(CourseSkillCoverage, CurriculumCourse).join(CurriculumCourse, CurriculumCourse.id == CourseSkillCoverage.course_id).where(CurriculumCourse.plan_id == plan_id, CourseSkillCoverage.skill_code.in_(skill_codes))).all()) if skill_codes else []
        by_skill: dict[str, list[tuple[CourseSkillCoverage, CurriculumCourse]]] = {code: [] for code in skill_codes}
        for mapping, course in mappings: by_skill.setdefault(mapping.skill_code, []).append((mapping, course))
        snapshots = {s.skill_code: s for s in self.db.execute(select(SkillDemandSnapshot).where(SkillDemandSnapshot.job_id == job_id, SkillDemandSnapshot.window_start.is_(None), SkillDemandSnapshot.skill_code.in_(skill_codes))).scalars()} if skill_codes else {}
        real_total = self.db.execute(select(func.count()).select_from(JobPosting).where(JobPosting.job_id == job_id, JobPosting.data_flag == DataFlag.REAL)).scalar_one()
        counts = {"covered": 0, "partial": 0, "uncovered": 0}
        for code in sorted(skill_codes):
            entries = by_skill.get(code, [])
            strength = max((m.coverage_strength for m, _ in entries), default=0)
            status = "covered" if strength >= 2 else "partial" if strength == 1 else "uncovered"
            counts[status] += 1
            snap = snapshots.get(code)
            evidence = [{"mapping_id": m.id, "course_id": c.id, "course_name": c.name, "chunk_id": m.source_chunk_id, "page": m.source_page, "quote": m.evidence_quote, "origin": m.origin} for m, c in entries]
            reliable_real_snapshot = bool(snap and snap.demo_posting_count == 0 and real_total)
            real_mentions = snap.posting_count if reliable_real_snapshot else 0
            self.db.add(SkillCoverage(analysis_id=analysis.id, skill_code=code, coverage_status=status, coverage_ratio=strength / 2, demand_frequency=snap.frequency if reliable_real_snapshot else None, posting_count=real_mentions, mastery_level=mastery[code], evidence=evidence))
        total = len(skill_codes)
        analysis.metrics = {**counts, "total": total, "covered_ratio": counts["covered"] / total if total else 0}
        analysis.status = "completed"
        self.db.flush()
        return analysis

    def analysis_payload(self, analysis: CurriculumAnalysis) -> dict[str, Any]:
        rows = list(self.db.execute(select(SkillCoverage).where(SkillCoverage.analysis_id == analysis.id).order_by(SkillCoverage.skill_code)).scalars())
        names = {s.skill_code: s.name_zh for s in self.db.execute(select(Skill).where(Skill.skill_code.in_([r.skill_code for r in rows]))).scalars()} if rows else {}
        courses = list(self.db.execute(select(CurriculumCourse).where(CurriculumCourse.plan_id == analysis.plan_id)).scalars())
        warnings: list[str] = []
        rich_courses = [course for course in courses if any(getattr(course, field) for field in COURSE_FIELDS)]
        if courses and not rich_courses:
            warnings.append("当前资料只提取到课程名称和学时，未提取到课程目标、教学内容或学习成果；覆盖判定会偏保守，建议继续上传对应课程大纲或课程信息。")
        if not rows:
            warnings.append("已审核能力图谱中暂无可分析的规范技能节点，请先完善能力图谱。")
        skills_payload = []
        for r in rows:
            course_items = [{
                "id": item.get("mapping_id") or self.db.execute(select(CourseSkillCoverage.id).where(CourseSkillCoverage.course_id == item.get("course_id"), CourseSkillCoverage.skill_code == r.skill_code)).scalar_one_or_none() or f"{r.skill_code}:{idx}",
                "course_id": item.get("course_id"),
                "course_name": item.get("course_name"),
                "skill_code": r.skill_code,
                "coverage_status": r.coverage_status,
                "origin": item.get("origin", "rule"),
                "teacher_confirmed": item.get("origin") == "teacher",
                "evidence_quote": item.get("quote", ""),
                "evidence": {"chunk_id": item.get("chunk_id"), "page": item.get("page"), "quote": item.get("quote")},
            } for idx, item in enumerate(r.evidence or [])]
            skills_payload.append({
                "skill_code": r.skill_code,
                "skill_name": names.get(r.skill_code),
                "coverage_status": r.coverage_status,
                "status": r.coverage_status,
                "coverage_ratio": r.coverage_ratio,
                "demand_frequency": r.demand_frequency,
                "posting_count": r.posting_count,
                "mastery_level": r.mastery_level,
                "graph_mastery": r.mastery_level,
                "evidence": r.evidence,
                "courses": course_items,
            })
        metrics = dict(analysis.metrics or {})
        # These values are derived from the current persisted rows rather than
        # the compact JSON snapshot written when the analysis was created.
        # Recompute them so legacy analyses (whose snapshot may contain zero)
        # stay accurate after a course or mapping edit.
        metrics["course_count"] = len(courses)
        metrics["mapped_skill_count"] = sum(1 for r in rows if r.coverage_status != "uncovered")
        metrics["covered_ratio"] = metrics.get("covered", 0) / max(metrics.get("total", len(rows)), 1)
        return {"id": analysis.id, "job_id": analysis.job_id, "plan_id": analysis.plan_id, "graph_id": analysis.graph_id, "snapshot_version": analysis.snapshot_version, "status": analysis.status, "metrics": metrics, "skills": skills_payload, "warnings": warnings}

    def generate_optimization(self, analysis_id: str) -> OptimizationRun:
        analysis = self.db.get(CurriculumAnalysis, analysis_id)
        if not analysis: raise NotFoundError(f"未找到课程对标分析：{analysis_id}")
        real_count = self.db.execute(select(func.count()).select_from(JobPosting).where(JobPosting.job_id == analysis.job_id, JobPosting.data_flag == DataFlag.REAL)).scalar_one()
        if real_count == 0:
            raise ValidationError("没有 REAL 岗位数据，不能生成正式优化建议。")
        coverages = list(self.db.execute(select(SkillCoverage).where(SkillCoverage.analysis_id == analysis.id)).scalars())
        skill_names = {s.skill_code: s.name_zh for s in self.db.execute(select(Skill).where(Skill.skill_code.in_([c.skill_code for c in coverages]))).scalars()} if coverages else {}
        gap_specs: list[dict[str, Any]] = []
        for row in coverages:
            gap_weight = {"uncovered": 1.0, "partial": 0.5, "covered": 0.0}[row.coverage_status]
            if gap_weight == 0 or row.demand_frequency is None:
                continue
            score = round(100 * gap_weight * (0.5 * row.demand_frequency + 0.2 * min(row.posting_count / 30, 1) + 0.3 * row.mastery_level / 4), 1)
            priority = "high" if score >= 60 else "medium" if score >= 30 else "low"
            name = skill_names.get(row.skill_code, row.skill_code)
            if row.coverage_status == "uncovered":
                action = "add_course" if priority == "high" else "add_teaching_content" if priority == "medium" else "adjust_course_objectives"
            elif priority == "high" and row.mastery_level >= 4:
                action = "add_project_practice"
            elif priority == "high":
                action = "add_practicum"
            elif priority == "medium" and row.posting_count >= 15:
                action = "adjust_hours"
            elif priority == "medium":
                action = "modify_course_content"
            else:
                action = "strengthen_competency"
            content = (f"建议在专业核心课或集中实训中补充{name}的教学、岗位情境任务和可评价成果。" if row.coverage_status == "uncovered" else f"现有课程已涉及{name}，建议补充实践任务、成果要求与评分 Rubric。")
            gap_specs.append({"row": row, "score": score, "priority": priority, "name": name, "action": action, "title": f"{'补充' if row.coverage_status == 'uncovered' else '强化'}「{name}」教学覆盖", "content": content})

        ai_drafts: dict[str, Any] = {}
        ai_run_id: str | None = None
        ai_confidence: float | None = None
        ai_note = "未配置真实 .env 中的 LLM_API_KEY，已使用确定性规则模板；统计、排序和证据不受影响。"
        if gap_specs and settings.llm_api_key:
            try:
                from app.agents.curriculum_optimization import CurriculumOptimizationAgent, CurriculumOptimizationInput
                envelope = CurriculumOptimizationAgent().run(self.db, CurriculumOptimizationInput(
                    job_id=analysis.job_id,
                    plan_id=analysis.plan_id,
                    gaps=[{
                        "skill_code": spec["row"].skill_code,
                        "skill_name": spec["name"],
                        "coverage_status": spec["row"].coverage_status,
                        "demand_frequency": spec["row"].demand_frequency,
                        "posting_count": spec["row"].posting_count,
                        "real_posting_total": real_count,
                        "mastery_level": spec["row"].mastery_level,
                        "priority": spec["priority"],
                        "priority_score": spec["score"],
                        "course_evidence": (spec["row"].evidence or [])[:5],
                    } for spec in gap_specs],
                ))
                drafts = envelope.result.suggestions
                expected = {spec["row"].skill_code for spec in gap_specs}
                if len(drafts) != len(expected) or {draft.skill_code for draft in drafts} != expected:
                    raise ValidationError("建议 Agent 返回的技能集合与确定性缺口不一致。")
                ai_drafts = {draft.skill_code: draft for draft in drafts}
                ai_run_id, ai_confidence = envelope.llm_run_id, envelope.confidence
                ai_note = envelope.reasoning_summary
            except Exception:
                ai_drafts = {}
                ai_run_id = None
                ai_note = "建议 Agent 调用或结构校验失败，已安全回退到确定性规则模板；统计、排序和证据不受影响。"
        for old in self.db.execute(select(OptimizationRun).where(OptimizationRun.job_id == analysis.job_id, OptimizationRun.plan_id == analysis.plan_id)).scalars():
            old.is_stale = True
        run = OptimizationRun(id=f"opt_{uuid.uuid4().hex}", analysis_id=analysis.id, job_id=analysis.job_id, plan_id=analysis.plan_id, graph_id=analysis.graph_id, snapshot_version=analysis.snapshot_version, low_sample=real_count < 30)
        self.db.add(run); self.db.flush()
        for spec in gap_specs:
            row = spec["row"]
            draft = ai_drafts.get(row.skill_code)
            suggestion = OptimizationSuggestion(id=f"sug_{uuid.uuid4().hex}", run_id=run.id, skill_code=row.skill_code, priority_score=spec["score"], priority=spec["priority"], action_type=draft.action_type if draft else spec["action"], title=draft.title if draft else spec["title"], content=draft.content if draft else spec["content"])
            self.db.add(suggestion); self.db.flush()
            self.db.add(SuggestionEvidence(suggestion_id=suggestion.id, evidence_type="ai_generation" if draft else "generation_note", source_id=ai_run_id or "deterministic_fallback", quote=draft.reason if draft else ai_note, extra={"ai_generated": bool(draft), "generation_run_id": ai_run_id, "confidence": ai_confidence}))
            snap = self.db.execute(select(SkillDemandSnapshot).where(SkillDemandSnapshot.job_id == analysis.job_id, SkillDemandSnapshot.skill_code == row.skill_code, SkillDemandSnapshot.window_start.is_(None))).scalar_one_or_none()
            if snap:
                quote = f"REAL 岗位 {row.posting_count}/{real_count} 条提及该技能，需求频率 {(row.demand_frequency or 0):.4f}。"
                self.db.add(SuggestionEvidence(suggestion_id=suggestion.id, evidence_type="market_snapshot", source_id=str(snap.id), quote=quote, extra={"posting_count": row.posting_count, "total_postings": real_count, "frequency": row.demand_frequency}))
                for posting, link in self.db.execute(select(JobPosting, JobPostingSkill).join(JobPostingSkill, JobPostingSkill.posting_id == JobPosting.id).where(JobPosting.job_id == analysis.job_id, JobPosting.data_flag == DataFlag.REAL, JobPostingSkill.skill_code == row.skill_code).limit(5)).all():
                    self.db.add(SuggestionEvidence(suggestion_id=suggestion.id, evidence_type="job_posting", source_id=posting.id, quote=link.evidence_span or posting.raw_text[:400], extra={"title": posting.title, "source_url": posting.source_url}))
            for ev in row.evidence:
                self.db.add(SuggestionEvidence(suggestion_id=suggestion.id, evidence_type="course_quote", source_id=ev["course_id"], quote=ev["quote"], page=ev.get("page"), extra=ev))
            node = self.db.execute(select(CompetencyNode).where(CompetencyNode.graph_id == analysis.graph_id, CompetencyNode.skill_code == row.skill_code).order_by(CompetencyNode.mastery_level.desc())).scalars().first()
            if node:
                quote = next((e.get("quote") for e in node.evidence if e.get("quote")), node.description or node.name)
                self.db.add(SuggestionEvidence(suggestion_id=suggestion.id, evidence_type="approved_graph", source_id=node.id, quote=quote, extra={"node_name": node.name, "mastery_level": row.mastery_level}))
        self.db.flush()
        return run

    def run_payload(self, run: OptimizationRun) -> dict[str, Any]:
        suggestions = list(self.db.execute(select(OptimizationSuggestion).where(OptimizationSuggestion.run_id == run.id).order_by(OptimizationSuggestion.priority_score.desc(), OptimizationSuggestion.skill_code)).scalars())
        names = {s.skill_code: s.name_zh for s in self.db.execute(select(Skill).where(Skill.skill_code.in_([s.skill_code for s in suggestions]))).scalars()} if suggestions else {}
        result = []
        for item in suggestions:
            coverage = self.db.execute(select(SkillCoverage).where(SkillCoverage.analysis_id == run.analysis_id, SkillCoverage.skill_code == item.skill_code)).scalar_one_or_none()
            evidence = list(self.db.execute(select(SuggestionEvidence).where(SuggestionEvidence.suggestion_id == item.id).order_by(SuggestionEvidence.id)).scalars())
            evidence_items = [{"type": e.evidence_type, "source_id": e.source_id, "quote": e.quote, "page": e.page, "extra": e.extra} for e in evidence]
            generation = next((e for e in evidence_items if e["type"] in {"ai_generation", "generation_note"}), None)
            job_evidence = [{"title": (e.get("extra") or {}).get("title", "岗位需求快照"), "quote": e["quote"], "source_url": (e.get("extra") or {}).get("source_url"), **e} for e in evidence_items if e["type"] in {"market_snapshot", "job_posting"}]
            course_evidence = [{"course_name": (e.get("extra") or {}).get("course_name", (e.get("extra") or {}).get("course_id", "课程")), "quote": e["quote"], "page": e.get("page"), "chunk_id": (e.get("extra") or {}).get("chunk_id"), **e} for e in evidence_items if e["type"] == "course_quote"]
            graph_evidence = next(({"node_name": (e.get("extra") or {}).get("node_name", "已审核能力节点"), "mastery_level": (e.get("extra") or {}).get("mastery_level"), **e} for e in evidence_items if e["type"] == "approved_graph"), None)
            coverage_status = coverage.coverage_status if coverage else "uncovered"
            demand_frequency = coverage.demand_frequency if coverage else None
            posting_count = coverage.posting_count if coverage else 0
            mastery_level = coverage.mastery_level if coverage else None
            affected_courses = sorted({e["course_name"] for e in course_evidence if e.get("course_name")})
            frequency_text = f"{demand_frequency * 100:.1f}%" if demand_frequency is not None else "无可靠 REAL 统计"
            reason = f"REAL 岗位中有 {posting_count} 条提及该技能（需求频率 {frequency_text}），图谱掌握度 {mastery_level or '—'}/4，当前课程为{'未覆盖' if coverage_status == 'uncovered' else '部分覆盖'}。"
            result.append({"id": item.id, "run_id": item.run_id, "skill_code": item.skill_code, "skill_name": names.get(item.skill_code), "priority_score": item.priority_score, "priority": item.priority, "action_type": item.action_type, "title": item.title, "content": item.content, "suggestion": item.content, "status": item.status, "state": item.status, "teacher_note": item.teacher_note, "coverage_status": coverage_status, "demand_frequency": demand_frequency, "posting_count": posting_count, "mastery_level": mastery_level, "affected_courses": affected_courses, "generation_reason": reason, "ai_generated": bool((generation or {}).get("extra", {}).get("ai_generated")), "generation_run_id": (generation or {}).get("extra", {}).get("generation_run_id"), "wording_reason": (generation or {}).get("quote"), "evidence": {
                "skill_code": item.skill_code,
                "job_evidence": job_evidence,
                "graph_evidence": graph_evidence,
                "course_evidence": course_evidence,
                "items": evidence_items,
            }})
        warnings = ["REAL 岗位样本少于 30 条，本结果为探索性分析。"] if run.low_sample else []
        if result and not any(item["ai_generated"] for item in result):
            warnings.append("建议文字使用确定性规则模板；配置真实 .env 后可由结构化 Agent 起草，统计、优先级和证据链保持不变。")
        return {"id": run.id, "analysis_id": run.analysis_id, "job_id": run.job_id, "plan_id": run.plan_id, "graph_id": run.graph_id, "status": "stale" if run.is_stale else "active", "is_stale": run.is_stale, "low_sample": run.low_sample, "snapshot_version": run.snapshot_version, "warnings": warnings, "suggestions": result}

    def export_docx(self, run: OptimizationRun) -> bytes:
        import docx
        payload = self.run_payload(run)
        plan = self.db.get(CurriculumPlan, run.plan_id)
        analysis = self.db.get(CurriculumAnalysis, run.analysis_id)
        coverage_rows = list(self.db.execute(select(SkillCoverage).where(SkillCoverage.analysis_id == run.analysis_id).order_by(SkillCoverage.demand_frequency.desc())).scalars())
        skill_names = {s.skill_code: s.name_zh for s in self.db.execute(select(Skill).where(Skill.skill_code.in_([row.skill_code for row in coverage_rows]))).scalars()} if coverage_rows else {}
        document = docx.Document()
        document.add_heading("人才培养方案优化建议报告", 0)
        document.add_paragraph(f"培养方案：{plan.name if plan else run.plan_id}")
        document.add_paragraph(f"岗位：{run.job_id}　能力图谱：{run.graph_id}")
        document.add_paragraph(f"岗位需求快照：{run.snapshot_version}")
        if run.low_sample: document.add_paragraph("注意：REAL 岗位少于 30 条，本报告为低样本探索性结果。")
        real_posting_count = self.db.scalar(
            select(func.count(JobPosting.id)).where(
                JobPosting.job_id == run.job_id,
                JobPosting.data_flag == DataFlag.REAL,
            )
        ) or 0
        document.add_heading("岗位需求概况", level=1)
        document.add_paragraph(
            f"本次分析使用 {real_posting_count} 条 REAL 岗位样本，"
            f"需求快照版本为 {run.snapshot_version}。"
        )
        document.add_heading("关键岗位能力", level=1)
        for row in coverage_rows[:10]:
            demand = f"{row.demand_frequency * 100:.1f}%" if row.demand_frequency is not None else "无可靠 REAL 统计"
            document.add_paragraph(f"{skill_names.get(row.skill_code, row.skill_code)}：{row.posting_count} 条 REAL 岗位，需求频率 {demand}，图谱掌握度 {row.mastery_level}/4", style="List Bullet")
        document.add_heading("课程覆盖情况", level=1)
        metrics = analysis.metrics if analysis else {}
        document.add_paragraph(f"能力总数：{metrics.get('total', len(coverage_rows))}；已覆盖：{metrics.get('covered', 0)}；部分覆盖：{metrics.get('partial', 0)}；未覆盖：{metrics.get('uncovered', 0)}；整体覆盖率：{metrics.get('covered_ratio', 0) * 100:.1f}%")
        coverage_labels = {"covered": "已覆盖", "partial": "部分覆盖", "uncovered": "未覆盖"}
        for row in coverage_rows:
            document.add_paragraph(f"{skill_names.get(row.skill_code, row.skill_code)}：{coverage_labels.get(row.coverage_status, row.coverage_status)}（覆盖比例 {row.coverage_ratio * 100:.0f}%）", style="List Bullet")
        document.add_heading("能力缺口", level=1)
        gap_rows = [row for row in coverage_rows if row.coverage_status != "covered"]
        if gap_rows:
            for row in gap_rows:
                demand = f"{row.demand_frequency * 100:.1f}%" if row.demand_frequency is not None else "无可靠 REAL 统计"
                document.add_paragraph(
                    f"{skill_names.get(row.skill_code, row.skill_code)}：{coverage_labels.get(row.coverage_status, row.coverage_status)}，"
                    f"岗位需求频率 {demand}。",
                    style="List Bullet",
                )
        else:
            document.add_paragraph("当前分析范围内未发现部分覆盖或未覆盖能力。")
        document.add_heading("培养方案优化建议", level=1)
        priority_labels = {"high": "高", "medium": "中", "low": "低"}
        state_labels = {"pending": "待确认", "adopted": "已采纳", "ignored": "已忽略"}
        action_labels = {
            "add_course": "新增课程", "adjust_course_objectives": "调整课程目标",
            "add_teaching_content": "增加教学内容", "add_practicum": "增加实训",
            "add_project_practice": "增加项目实践", "adjust_hours": "调整课时",
            "strengthen_competency": "强化能力", "modify_course_content": "修改课程内容",
        }
        for item in payload["suggestions"]:
            document.add_heading(f"[{priority_labels.get(item['priority'], item['priority'])} {item['priority_score']:.1f}] {item['title']}", level=2)
            document.add_paragraph(item["content"])
            document.add_paragraph(f"建议类型：{action_labels.get(item['action_type'], item['action_type'])}；当前覆盖：{coverage_labels.get(item['coverage_status'], item['coverage_status'])}；涉及课程：{'、'.join(item['affected_courses']) if item['affected_courses'] else '暂无对应课程'}")
            document.add_paragraph(f"生成依据：{item['generation_reason']}")
            document.add_paragraph(f"教师决策：{state_labels.get(item['status'], item['status'])}")
            if item["teacher_note"]: document.add_paragraph(f"教师备注：{item['teacher_note']}")
            document.add_heading("证据", level=2)
            evidence = item["evidence"]
            evidence_rows = evidence.get("items", []) if isinstance(evidence, dict) else evidence
            for ev in evidence_rows:
                document.add_paragraph(f"{ev.get('type', 'evidence')} / {ev.get('source_id', '')}：{ev.get('quote', '')}", style="List Bullet")
        output = BytesIO(); document.save(output); return output.getvalue()
