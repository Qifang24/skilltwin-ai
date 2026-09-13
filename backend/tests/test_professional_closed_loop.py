"""End-to-end regression for the professional-construction closed loop."""

from __future__ import annotations

import json
from io import BytesIO

import docx
from pypdf import PdfWriter
from sqlalchemy import func, select

from app.core.enums import GraphStatus, NodeType, Provenance, SkillCategory, SkillStatus
from app.models.competency import CompetencyGraph, CompetencyNode
from app.models.curriculum import CourseSkillCoverage, CurriculumAnalysis, CurriculumCourse, OptimizationSuggestion
from app.models.job_market import JobImportRow, JobPosting
from app.models.ontology import Skill
from app.services.curriculum_management_service import _table_courses


def _seed_skills(db) -> None:
    db.add_all([
        Skill(skill_code="prog.python", name_zh="Python", name_en="Python", category=SkillCategory.PROGRAMMING, aliases=["Python编程"], status=SkillStatus.ACTIVE, provenance=Provenance.HUMAN_AUTHORED),
        Skill(skill_code="quality.data", name_zh="数据质量", name_en="Data Quality", category=SkillCategory.QUALITY, aliases=["质量检查"], status=SkillStatus.ACTIVE, provenance=Provenance.HUMAN_AUTHORED),
    ])
    db.commit()


def _import_real_job(client) -> str:
    jd = "负责 Python 数据处理与数据质量检查，联系手机 13812345678，确保标注结果准确。"
    payload = [
        {"id": "source-1", "title": "数据工程师", "company_name": "示例企业", "raw_text": jd, "source_name": "公开招聘站", "source_url": "https://example.test/jobs/1", "posted_at": "2026-09-01"},
        {"id": "source-2", "title": "重复岗位", "company_name": "示例企业", "raw_text": jd, "source_name": "公开招聘站", "source_url": "https://example.test/jobs/2", "posted_at": "2026-09-02"},
    ]
    response = client.post(
        "/api/v1/job-imports",
        data={"job_id": "data_engineer", "job_name": "数据工程师"},
        files={"file": ("jobs.json", json.dumps(payload, ensure_ascii=False).encode(), "application/json")},
    )
    assert response.status_code == 201, response.text
    batch = response.json()
    assert batch["total_rows"] == 2
    assert [row["status"] for row in batch["rows"]] == ["valid", "duplicate"]
    return batch["id"]


def test_staged_job_import_is_scrubbed_deduplicated_and_idempotent(client, db) -> None:
    _seed_skills(db)
    batch_id = _import_real_job(client)

    assert db.scalar(select(func.count()).select_from(JobPosting)) == 0
    staged = db.execute(select(JobImportRow).where(JobImportRow.batch_id == batch_id).order_by(JobImportRow.row_number)).scalars().all()
    assert "13812345678" not in json.dumps([row.raw_data for row in staged], ensure_ascii=False)

    first = client.post(f"/api/v1/job-imports/{batch_id}/confirm")
    assert first.status_code == 200, first.text
    assert first.json()["success_count"] == 1
    assert first.json()["duplicate_count"] == 1
    second = client.post(f"/api/v1/job-imports/{batch_id}/confirm")
    assert second.status_code == 200
    assert second.json()["success_count"] == 1
    assert db.scalar(select(func.count()).select_from(JobPosting)) == 1
    posting = db.execute(select(JobPosting)).scalar_one()
    assert posting.pii_scrubbed is True
    assert "13812345678" not in posting.raw_text
    assert client.get("/api/v1/job-market/data_engineer/postings").json()[0]["id"] == posting.id

    # Historical duplicates are detected during preview, before confirmation.
    historical = client.post(
        "/api/v1/job-imports", data={"job_id": "data_engineer", "job_name": "数据工程师"},
        files={"file": ("again.json", json.dumps([{"title": "再次抓取", "raw_text": "负责 Python 数据处理与数据质量检查，联系手机 13812345678，确保标注结果准确。", "source_name": "公开招聘站"}], ensure_ascii=False).encode(), "application/json")},
    )
    assert historical.status_code == 201
    assert historical.json()["rows"][0]["status"] == "duplicate"


def test_curriculum_analysis_teacher_recalculation_optimization_and_docx(client, db) -> None:
    _seed_skills(db)
    batch_id = _import_real_job(client)
    assert client.post(f"/api/v1/job-imports/{batch_id}/confirm").status_code == 200

    curriculum = {"courses": [{
        "id": "course-python", "name": "Python 数据处理", "course_code": "AI101", "total_hours": 48,
        "objectives": "掌握 Python 编程基础。", "teaching_content": "使用 Python 完成数据清洗。",
        "learning_outcomes": "能够提交可运行的数据处理程序。",
    }], "coverage": [{
        "course_id": "course-python", "skill_code": "prog.python", "coverage_strength": 1,
        "evidence_quote": "使用 Python 完成数据清洗。", "source_chunk_id": "json:1",
        "teacher_confirmed": True, "edited_by": "teacher",
    }]}
    uploaded = client.post(
        "/api/v1/curriculum/imports",
        data={"source_name": "2026 人才培养方案", "license_note": "校内授权材料"},
        files={"file": ("plan.json", json.dumps(curriculum, ensure_ascii=False).encode(), "application/json")},
    )
    assert uploaded.status_code == 201, uploaded.text
    import_id = uploaded.json()["id"]
    assert uploaded.json()["status"] == "parsed"
    assert db.scalar(select(func.count()).select_from(CurriculumCourse)) == 0

    confirmed = client.post(f"/api/v1/curriculum/imports/{import_id}/confirm", json={
        "plan_id": "plan_ai_2026", "name": "人工智能技术应用 2026", "profession": "人工智能技术应用", "version": "2026", "is_partial": False,
    })
    assert confirmed.status_code == 200, confirmed.text
    course = db.execute(select(CurriculumCourse).where(CurriculumCourse.plan_id == "plan_ai_2026")).scalar_one()
    assert course.id == "course-python"
    imported_mapping = db.execute(select(CourseSkillCoverage).where(CourseSkillCoverage.course_id == course.id, CourseSkillCoverage.skill_code == "prog.python")).scalar_one()
    assert imported_mapping.evidence_quote == "使用 Python 完成数据清洗。"
    assert imported_mapping.teacher_confirmed is True

    graph = CompetencyGraph(id="graph_data_v1", job_id="data_engineer", version=1, status=GraphStatus.APPROVED, title="数据工程师能力图谱", selected_skill_codes=["prog.python", "quality.data"], approved_by="teacher")
    db.add(graph)
    db.add_all([
        CompetencyNode(id="graph_data_v1.python", graph_id=graph.id, node_type=NodeType.SKILL_POINT, name="Python 编程", skill_code="prog.python", mastery_level=3, evidence=[{"quote": "负责 Python 数据处理"}], ai_generated=False),
        CompetencyNode(id="graph_data_v1.quality", graph_id=graph.id, node_type=NodeType.SKILL_POINT, name="数据质量控制", skill_code="quality.data", mastery_level=4, evidence=[{"quote": "负责数据质量检查"}], ai_generated=False),
    ])
    db.commit()

    analysis_response = client.post("/api/v1/curriculum/analyses", json={"job_id": "data_engineer", "plan_id": "plan_ai_2026", "graph_id": graph.id})
    assert analysis_response.status_code == 201, analysis_response.text
    analysis = analysis_response.json()
    by_code = {row["skill_code"]: row for row in analysis["skills"]}
    assert by_code["prog.python"]["status"] == "partial"
    assert by_code["quality.data"]["status"] == "uncovered"
    assert by_code["prog.python"]["courses"][0]["id"] == imported_mapping.id
    assert by_code["prog.python"]["courses"][0]["evidence_quote"] == "使用 Python 完成数据清洗。"
    assert analysis["metrics"]["course_count"] == 1
    repeated = client.post("/api/v1/curriculum/analyses", json={"job_id": "data_engineer", "plan_id": "plan_ai_2026", "graph_id": graph.id})
    assert repeated.status_code == 201, repeated.text
    assert repeated.json()["id"] == analysis["id"]
    assert db.scalar(select(func.count()).select_from(CurriculumAnalysis)) == 1

    mapping = db.execute(select(CourseSkillCoverage).where(CourseSkillCoverage.course_id == course.id, CourseSkillCoverage.skill_code == "prog.python")).scalar_one()
    changed = client.patch(f"/api/v1/curriculum/skill-mappings/{mapping.id}", json={
        "skill_code": "prog.python", "coverage_status": "covered", "evidence_quote": "使用 Python 完成数据清洗。",
        "source_chunk_id": course.source_chunk_id, "teacher_confirmed": True, "edited_by": "teacher", "reason": "已核查课程标准",
    })
    assert changed.status_code == 200, changed.text
    refreshed = client.get("/api/v1/curriculum/analyses", params={"job_id": "data_engineer", "plan_id": "plan_ai_2026", "graph_id": graph.id})
    assert refreshed.status_code == 200
    assert {row["skill_code"]: row["status"] for row in refreshed.json()["skills"]}["prog.python"] == "covered"
    db.expire_all()
    assert db.get(CourseSkillCoverage, mapping.id).teacher_confirmed is True

    generated = client.post("/api/v1/curriculum/optimization-runs", json={"job_id": "data_engineer", "plan_id": "plan_ai_2026", "graph_id": graph.id})
    assert generated.status_code == 201, generated.text
    run = generated.json()
    assert run["status"] == "active"
    assert run["low_sample"] is True
    assert [item["skill_code"] for item in run["suggestions"]] == ["quality.data"]
    suggestion = run["suggestions"][0]
    assert suggestion["priority"] == "high"
    assert suggestion["coverage_status"] == "uncovered"
    assert suggestion["posting_count"] == 1
    assert "REAL 岗位" in suggestion["generation_reason"]
    assert suggestion["evidence"]["job_evidence"]
    assert suggestion["evidence"]["graph_evidence"]

    adopted = client.patch(f"/api/v1/curriculum/optimization-suggestions/{suggestion['id']}", json={"state": "adopted", "action_type": "add_project_practice", "teacher_note": "纳入下一版课程"})
    assert adopted.status_code == 200, adopted.text
    assert adopted.json()["state"] == "adopted"
    assert adopted.json()["action_type"] == "add_project_practice"
    db.expire_all()
    assert db.get(OptimizationSuggestion, suggestion["id"]).status == "adopted"

    report = client.get(f"/api/v1/curriculum/optimization-runs/{run['id']}/report.docx")
    assert report.status_code == 200
    parsed = docx.Document(BytesIO(report.content))
    report_text = "\n".join(paragraph.text for paragraph in parsed.paragraphs)
    assert "人才培养方案优化建议报告" in report_text
    assert "岗位需求概况" in report_text
    assert "关键岗位能力" in report_text
    assert "课程覆盖情况" in report_text
    assert "能力缺口" in report_text
    assert "培养方案优化建议" in report_text
    assert "纳入下一版课程" in report_text
    assert "数据质量" in report_text

    rebuilt = client.post("/api/v1/job-market/data_engineer/analyze")
    assert rebuilt.status_code == 200, rebuilt.text
    stale = client.get("/api/v1/curriculum/optimization-runs/latest", params={
        "job_id": "data_engineer", "plan_id": "plan_ai_2026", "graph_id": graph.id,
    })
    assert stale.status_code == 200
    assert stale.json()["status"] == "stale"
    db.expire_all()
    assert db.get(OptimizationSuggestion, suggestion["id"]).status == "adopted"


def test_mapping_rejects_non_contiguous_quote(client, db) -> None:
    _seed_skills(db)
    from app.models.curriculum import CurriculumPlan
    db.add(CurriculumPlan(id="p", name="方案", source_name="来源"))
    db.flush()
    db.add(CurriculumCourse(id="c", plan_id="p", name="课程", source_chunk_id="manual", source_quote="课程只包含 Python 原文。", teaching_content="Python 实践"))
    db.commit()
    response = client.post("/api/v1/curriculum/courses/c/skills", json={"skill_code": "prog.python", "coverage_status": "covered", "evidence_quote": "不存在的改写", "teacher_confirmed": True, "edited_by": "teacher"})
    assert response.status_code == 422
    empty = client.post("/api/v1/curriculum/courses/c/skills", json={"skill_code": "prog.python", "coverage_status": "covered", "evidence_quote": "", "teacher_confirmed": True, "edited_by": "teacher"})
    assert empty.status_code == 422


def test_supported_upload_formats_and_scanned_pdf_message(client, db) -> None:
    # GB18030 is a required job-import encoding and a missing date is warning-only.
    csv_text = "岗位名称,岗位描述,来源\n数据治理工程师,负责数据治理与规范建设,公开招聘站\n"
    job = client.post(
        "/api/v1/job-imports", data={"job_id": "governance", "job_name": "数据治理"},
        files={"file": ("jobs.csv", csv_text.encode("gb18030"), "text/csv")},
    )
    assert job.status_code == 201, job.text
    assert job.json()["rows"][0]["status"] == "valid"
    assert any("趋势" in warning for warning in job.json()["rows"][0]["warnings"])

    text = client.post(
        "/api/v1/curriculum/imports", data={"source_name": "文本方案"},
        files={"file": ("plan.txt", "课程名称：数据治理\n课程目标：掌握治理规范\n教学内容：数据治理流程".encode(), "text/plain")},
    )
    assert text.status_code == 201
    assert text.json()["status"] == "parsed"
    assert text.json()["courses"][0]["data"]["name"] == "数据治理"

    document = docx.Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "课程名称", "质量管理实训"
    table.cell(1, 0).text, table.cell(1, 1).text = "教学内容", "数据质量检查"
    docx_bytes = BytesIO(); document.save(docx_bytes)
    word = client.post(
        "/api/v1/curriculum/imports", data={"source_name": "Word 方案"},
        files={"file": ("plan.docx", docx_bytes.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert word.status_code == 201, word.text
    assert word.json()["status"] == "parsed"
    assert word.json()["courses"][0]["data"]["name"] == "质量管理实训"

    writer = PdfWriter(); writer.add_blank_page(width=100, height=100)
    pdf_bytes = BytesIO(); writer.write(pdf_bytes)
    scanned = client.post(
        "/api/v1/curriculum/imports", data={"source_name": "扫描方案"},
        files={"file": ("scan.pdf", pdf_bytes.getvalue(), "application/pdf")},
    )
    assert scanned.status_code == 201
    assert scanned.json()["status"] == "failed"
    assert "扫描" in scanned.json()["error_message"] or "OCR" in scanned.json()["error_message"]


def test_curriculum_table_parser_extracts_pdf_style_course_rows() -> None:
    text = """课程代码 课程名称 学分 总学时
100081063
大数据分析技术
Big Data Analytics
2 32 24 8 32
100081062
大数据处理技术
Big Data Processing Technology
2 32 24 8 32
100081999
大数据分析技术
Big Data Analytics (duplicate table row)
2 32 24 8 32
"""
    rows = _table_courses(text)
    assert [(row[0]["course_code"], row[0]["name"], row[0]["total_hours"]) for row in rows] == [
        ("100081063", "大数据分析技术", 32),
        ("100081062", "大数据处理技术", 32),
    ]


def test_confirmation_skips_duplicate_legacy_course_names(client, db) -> None:
    uploaded = client.post(
        "/api/v1/curriculum/imports",
        data={"source_name": "重复课程方案"},
        files={"file": ("duplicate.json", json.dumps({"courses": [{"name": "重复课程", "total_hours": 32}, {"name": "重复课程", "total_hours": 48}]}, ensure_ascii=False).encode(), "application/json")},
    )
    assert uploaded.status_code == 201
    batch_id = uploaded.json()["id"]
    confirmed = client.post(f"/api/v1/curriculum/imports/{batch_id}/confirm", json={"plan_id": "plan_duplicate_legacy", "name": "重复课程方案"})
    assert confirmed.status_code == 200, confirmed.text
    assert db.scalar(select(func.count()).select_from(CurriculumCourse).where(CurriculumCourse.plan_id == "plan_duplicate_legacy")) == 1


def test_curriculum_plan_delete_removes_plan_and_courses(client, db) -> None:
    from app.models.curriculum import CurriculumPlan

    db.add(CurriculumPlan(id="plan_delete_test", name="待删除方案", source_name="测试"))
    db.commit()
    response = client.delete("/api/v1/curriculum/plans/plan_delete_test")
    assert response.status_code == 204
    assert db.get(CurriculumPlan, "plan_delete_test") is None
